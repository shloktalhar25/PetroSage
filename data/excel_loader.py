# -*- coding: utf-8 -*-
"""
data/excel_loader.py - Document source connector for Excel files (.xlsx, .xlsm).

For each sheet, the header row is found heuristically (first row with at
least two non-empty cells - this skips title rows and blank spacer rows),
then every subsequent non-empty row becomes one chunk-shaped block.
"""

from pathlib import Path
from typing import List, Dict, Any

# pyrefly: ignore [import-error, missing-import]
import openpyxl


def _find_header_row(rows: List[tuple]) -> int:
    """Return the index (0-based) of the first row with >= 2 non-empty cells."""
    for i, row in enumerate(rows):
        non_empty = sum(1 for v in row if v is not None and str(v).strip() != "")
        if non_empty >= 2:
            return i
    return 0


# Norwegian Offshore Directorate headers -> English, so English questions match.
_HEADER_EN = {
    "Selskapsnavn": "Company name",
    "Område": "Area",
    "Godkjent dato": "Approved date",
    "Feltnavn": "Field name",
    "Funnnavn": "Discovery name",
    "Operatør": "Operator",
    "Status": "Status",
    "#licensees": "Number of production licences held (as licensee)",
    "#operatorships": "Number of licences operated",
    "#operatorships fields": "Number of fields operated",
    "#operatorships discoveries": "Number of discoveries operated",
}


def _dataset_label(path: Path, sheet: str, title: str) -> str:
    """Human label prepended to every row so retrieval knows what the row is about."""
    name = path.stem.replace("_", " ")
    extra = [t for t in (title, sheet) if t and t.lower() not in name.lower()]
    return " - ".join([name] + extra[:1])


def load_excel(path: Path) -> List[Dict[str, Any]]:
    """Read an Excel workbook, returning one block per populated data row."""
    blocks: List[Dict[str, Any]] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        header_idx = _find_header_row(rows)
        headers = [_HEADER_EN.get(str(h).strip(), h) if h is not None else None
                   for h in rows[header_idx]]
        if headers and headers[0] is None:  # unlabeled first column = row label (e.g. year/month)
            headers[0] = "Period"
        title = next((str(c) for r in rows[:header_idx] for c in r if c), "")
        label = _dataset_label(path, sheet_name, title)

        for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if all(v is None for v in row):
                continue
            parts = [
                f"{h}: {v}"
                for h, v in zip(headers, row)
                if h is not None and v is not None and str(v).strip() != ""
            ]
            if not parts:
                continue
            blocks.append({
                "text": f"{label}: " + ", ".join(parts),
                "sheet": sheet_name,
                "row": i,
            })

    wb.close()
    return blocks
